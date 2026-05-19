import sys
import os
import subprocess
import traceback
import datetime
import tempfile
import threading
import asyncio
import queue
import ctypes
import uuid
import json
import time
import socket
import webbrowser
import platform
from functools import wraps

# ================== 隐藏控制台（Windows） ==================
if platform.system() == "Windows":
    try:
        ctypes.windll.user32.ShowWindow(ctypes.windll.kernel32.GetConsoleWindow(), 0)
    except:
        pass

# ================== 切换到脚本所在目录 ==================
if getattr(sys, 'frozen', False):
    BASE_DIR = os.path.dirname(sys.executable)
else:
    BASE_DIR = os.path.dirname(os.path.abspath(__file__))
os.chdir(BASE_DIR)

# ================== 调试日志 ==================
LOG_FILE = os.path.join(BASE_DIR, "debug.log")
def debug(msg):
    try:
        with open(LOG_FILE, "a", encoding="utf-8") as f:
            f.write(f"[{datetime.datetime.now().strftime('%H:%M:%S')}] {msg}\n")
    except:
        pass

debug("===== 班呼ClassCall 启动 =====")

def log_exception(exc_type, exc_value, exc_tb):
    msg = ''.join(traceback.format_exception(exc_type, exc_value, exc_tb))
    debug(f"未捕获异常:\n{msg}")
sys.excepthook = log_exception

# ================== 可调参数 ==================
DISPLAY_FONT_FAMILY = "微软雅黑"
DISPLAY_FONT_SIZE   = 48
DISPLAY_FONT_COLOR  = "white"

# ================== 系统主题检测 ==================
def is_windows_dark_mode():
    if platform.system() != "Windows":
        return False
    try:
        import winreg
        key = winreg.OpenKey(winreg.HKEY_CURRENT_USER,
                             r"Software\Microsoft\Windows\CurrentVersion\Themes\Personalize")
        value, _ = winreg.QueryValueEx(key, "AppsUseLightTheme")
        winreg.CloseKey(key)
        return value == 0
    except:
        return False

def get_system_bg_color():
    return "#1e1e1e" if is_windows_dark_mode() else "#f0f0f0"

# ================== 依赖安装 ==================
def install(pkg):
    cmds = [
        [sys.executable, "-m", "pip", "install", "--user", pkg],
        [sys.executable, "-m", "pip", "install", pkg],
    ]
    for cmd in cmds:
        try:
            subprocess.check_call(cmd, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
            return True
        except:
            continue
    return False

def safe_import(module, pip_name):
    try:
        __import__(module)
        debug(f"模块 {module} 导入成功")
        return True
    except ImportError:
        debug(f"模块 {module} 缺失，尝试安装 {pip_name}")
        if install(pip_name):
            debug(f"安装 {pip_name} 成功")
            try:
                __import__(module)
                return True
            except:
                pass
        debug(f"安装 {pip_name} 失败")
        return False

# 核心依赖
if not safe_import("flask", "flask"):
    sys.exit(1)
if not safe_import("qrcode", "qrcode"):
    sys.exit(1)
if not safe_import("PIL", "Pillow"):
    sys.exit(1)
EDGE_AVAILABLE = safe_import("edge_tts", "edge-tts")
OPENPYXL_AVAILABLE = safe_import("openpyxl", "openpyxl")
TRAY_AVAILABLE = safe_import("pystray", "pystray")
if not TRAY_AVAILABLE:
    debug("警告：pystray未安装，将无法使用系统托盘功能")

import tkinter as tk
from tkinter import ttk, scrolledtext, simpledialog, messagebox
from flask import Flask, request, render_template_string, jsonify, session, redirect, url_for
import qrcode
import pystray
from PIL import Image, ImageDraw

# ================== 配置管理 ==================
CONFIG_FILE = os.path.join(BASE_DIR, "config.json")
def load_config():
    default = {
        "bg_color": get_system_bg_color(),
        "sender_name": "班呼ClassCall",
        "display_mode": "classic",
        "speech_enabled": True,
        "banner_bg_color": "#2c3e50",
        "banner_text_color": "white",
        "banner_height": 60,
        "banner_font_size": 16,
        "banner_scroll_speed": 100,
        "classic_alpha": 1.0,
        "banner_alpha": 1.0,
        "text_color": DISPLAY_FONT_COLOR,
        "popup_bg_color": "#2c3e50",
        "popup_text_color": "white",
        "popup_font_size": 12,
        "popup_max_height": 300,
        "popup_auto_scroll": False,
        "popup_vertical_scroll_speed": 50,
        "web_password": "123456",
        "exit_password": "123456"
    }
    try:
        with open(CONFIG_FILE, "r", encoding="utf-8") as f:
            data = json.load(f)
            for k, v in default.items():
                if k not in data:
                    data[k] = v
            return data
    except:
        return default

def save_config(data):
    try:
        with open(CONFIG_FILE, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
    except:
        pass

config = load_config()
bg_color = config["bg_color"]
sender_name = config["sender_name"]
display_mode = config["display_mode"]
speech_enabled = config["speech_enabled"]
banner_bg_color = config["banner_bg_color"]
banner_text_color = config["banner_text_color"]
banner_height = config["banner_height"]
banner_font_size = config["banner_font_size"]
banner_scroll_speed = config.get("banner_scroll_speed", 100)
classic_alpha = config["classic_alpha"]
banner_alpha = config["banner_alpha"]
text_color = config.get("text_color", DISPLAY_FONT_COLOR)
popup_bg_color = config.get("popup_bg_color", "#2c3e50")
popup_text_color = config.get("popup_text_color", "white")
popup_font_size = config.get("popup_font_size", 12)
popup_max_height = config.get("popup_max_height", 300)
popup_auto_scroll = config.get("popup_auto_scroll", False)
popup_vertical_scroll_speed = config.get("popup_vertical_scroll_speed", 50)
WEB_PASSWORD = config.get("web_password", "123456")
EXIT_PASSWORD = config.get("exit_password", "123456")

# ================== 全局变量 ==================
stop_flag = threading.Event()
queue_mgr = None
classic_window = None
banner_window = None
popup_notification = None
root = None
main_ui_window = None
tray_icon = None
last_user_action_time = None      # 用户主动操作时间（用于在线状态）

# ================== 动画函数（保持不变） ==================
def set_window_alpha(window, alpha):
    try:
        window.attributes('-alpha', alpha)
    except:
        pass

def fade_in(window, target_alpha=1.0, duration=200, step=20):
    window.attributes('-alpha', 0.0)
    window.deiconify()
    steps = int(duration / step)
    delta = target_alpha / steps
    def animate(alpha=0.0):
        alpha += delta
        if alpha >= target_alpha:
            window.attributes('-alpha', target_alpha)
            return
        window.attributes('-alpha', alpha)
        window.after(step, lambda: animate(alpha))
    animate()

def fade_out(window, duration=200, step=20, on_complete=None):
    steps = int(duration / step)
    delta = 1.0 / steps
    def animate(alpha=1.0):
        alpha -= delta
        if alpha <= 0.0:
            window.attributes('-alpha', 0.0)
            window.withdraw()
            if on_complete:
                on_complete()
            return
        window.attributes('-alpha', alpha)
        window.after(step, lambda: animate(alpha))
    animate()

def slide_in_banner(window, start_y=-100, end_y=10, duration=300, target_alpha=1.0):
    window.geometry(f"+{window.winfo_x()}+{start_y}")
    window.attributes('-alpha', 0.0)
    window.deiconify()
    steps = 20
    step_time = duration // steps
    y_step = (end_y - start_y) / steps
    alpha_step = target_alpha / steps
    def animate(step=0, y=start_y, alpha=0.0):
        if step >= steps:
            window.geometry(f"+{window.winfo_x()}+{end_y}")
            window.attributes('-alpha', target_alpha)
            return
        y += y_step
        alpha += alpha_step
        window.geometry(f"+{window.winfo_x()}+{int(y)}")
        window.attributes('-alpha', alpha)
        window.after(step_time, lambda: animate(step+1, y, alpha))
    animate()

def slide_out_banner(window, end_y=-100, duration=300, on_complete=None):
    start_y = int(window.geometry().split('+')[2])
    steps = 20
    step_time = duration // steps
    y_step = (end_y - start_y) / steps
    alpha_step = 1.0 / steps
    def animate(step=0, y=start_y, alpha=1.0):
        if step >= steps:
            window.withdraw()
            if on_complete:
                on_complete()
            return
        y += y_step
        alpha -= alpha_step
        if alpha < 0: alpha = 0
        window.geometry(f"+{window.winfo_x()}+{int(y)}")
        window.attributes('-alpha', alpha)
        window.after(step_time, lambda: animate(step+1, y, alpha))
    animate()

# ================== 传统窗口（精简，保持原有功能） ==================
class ClassicWindow:
    def __init__(self, root, device_code, bg_color, sender_name):
        self.root = root
        self.root.title(f"班呼ClassCall - 叫号显示 [设备: {device_code}]")
        self.root.geometry("600x350+100+100")
        self.root.minsize(300, 150)
        self.root.attributes("-topmost", True)
        self.root.configure(bg=bg_color)
        self.root.attributes('-alpha', classic_alpha)
        self.root.withdraw()

        try:
            style = ttk.Style()
            style.theme_use('clam')
            style.configure("TButton", font=("微软雅黑", 10), padding=6, borderwidth=0)
            style.map("TButton", background=[('active', '#4a4a4a')])
        except:
            pass

        self.root.bind('<Escape>', self.escape_handler)
        self.root.bind('<F11>', lambda e: self.toggle_fullscreen())
        self.root.protocol("WM_DELETE_WINDOW", self.hide_window)
        self.root.bind('<Configure>', self.on_resize)

        self.font_family = DISPLAY_FONT_FAMILY
        self.font_size = DISPLAY_FONT_SIZE
        self.font_color = text_color

        self.sender_label = tk.Label(root, text=f"发件人: {sender_name}",
                                     font=("微软雅黑", 16, "bold"), fg="gray", bg=bg_color, anchor=tk.W)
        self.sender_label.pack(fill=tk.X, padx=20, pady=(12, 0))

        self.label = tk.Label(root, text="等待叫号...",
                             font=(self.font_family, self.font_size, "bold"),
                             fg=self.font_color, bg=bg_color, anchor=tk.CENTER, justify=tk.CENTER)
        self.label.pack(expand=True, fill=tk.BOTH, padx=20, pady=(0, 20))

        btn_frame = tk.Frame(root, bg=bg_color)
        btn_frame.pack(pady=5)
        ttk.Button(btn_frame, text="关闭窗口", command=self.hide_window).pack(side=tk.LEFT, padx=5)
        self.green_btn = ttk.Button(btn_frame, text="绿幕模式", command=self.toggle_green)
        self.green_btn.pack(side=tk.LEFT, padx=5)
        self.full_btn = ttk.Button(btn_frame, text="全屏", command=self.toggle_fullscreen)
        self.full_btn.pack(side=tk.LEFT, padx=5)

        self.green_mode = False
        self.fullscreen = False

        self.status_bar = tk.Label(root, text="队列: 0 | 状态: 空闲 | 模式: 手动",
                                   bg="#333", fg="white", font=("微软雅黑", 10), anchor=tk.W)
        self.status_bar.pack(fill=tk.X, side=tk.BOTTOM)

        self.text_queue = queue.Queue()
        self.hide_after_id = None
        self.update_wraplength()
        self.update_display()

    def set_alpha(self, alpha):
        set_window_alpha(self.root, alpha)

    def on_resize(self, event=None):
        self.update_wraplength()

    def update_wraplength(self):
        try:
            w = self.label.winfo_width()
            if w > 20:
                self.label.configure(wraplength=w - 20)
        except:
            pass

    def hide_window(self, animated=True):
        if self.hide_after_id:
            self.root.after_cancel(self.hide_after_id)
        if animated and self.root.winfo_viewable():
            fade_out(self.root, duration=200, on_complete=lambda: None)
        else:
            self.root.withdraw()

    def show_window_animated(self):
        if not self.root.winfo_viewable():
            fade_in(self.root, target_alpha=classic_alpha, duration=200)

    def schedule_hide(self, ms=1000, animated=True):
        if self.hide_after_id:
            self.root.after_cancel(self.hide_after_id)
        self.hide_after_id = self.root.after(ms, lambda: self.hide_window(animated))

    def escape_handler(self, e):
        if self.fullscreen:
            self.toggle_fullscreen()
        elif self.green_mode:
            self.toggle_green()

    def toggle_fullscreen(self):
        self.fullscreen = not self.fullscreen
        self.root.attributes('-fullscreen', self.fullscreen)
        self.full_btn.config(text="窗口" if self.fullscreen else "全屏")

    def toggle_green(self):
        self.green_mode = not self.green_mode
        bg = "#00FF00" if self.green_mode else bg_color
        self.green_btn.config(text="恢复黑色" if self.green_mode else "绿幕模式")
        self._apply_bg(bg)

    def _apply_bg(self, bg):
        try:
            self.root.configure(bg=bg)
            self.label.config(bg=bg)
            self.sender_label.config(bg=bg)
            for child in self.root.winfo_children():
                if isinstance(child, tk.Frame):
                    child.config(bg=bg)
            self.status_bar.config(bg="#008800" if self.green_mode else "#333")
        except:
            pass

    def update_sender(self, name):
        try:
            self.sender_label.config(text=f"发件人: {name}")
        except:
            pass

    def update_font(self, family=None, size=None, color=None):
        if family:
            self.font_family = family
        if size:
            self.font_size = size
        if color:
            self.font_color = color
        try:
            self.label.config(font=(self.font_family, self.font_size, "bold"), fg=self.font_color)
            self.update_wraplength()
        except:
            pass

    def update_status(self, queue_count, playing_status, mode):
        try:
            self.status_bar.config(text=f"队列: {queue_count} | 状态: {playing_status} | 模式: {mode}")
        except:
            pass

    def update_display(self):
        try:
            while True:
                text = self.text_queue.get_nowait()
                if self.hide_after_id:
                    self.root.after_cancel(self.hide_after_id)
                self.label.config(text=text)
                self.show_window_animated()
                self.update_wraplength()
        except queue.Empty:
            pass
        self.root.after(500, self.update_display)

    def show_text(self, text):
        self.text_queue.put(text)

    def clear_text(self):
        self.text_queue.queue.clear()
        self.label.config(text="等待叫号...")
        self.root.after(100, self.update_wraplength)

# ================== 横幅窗口 ==================
class BannerWindow:
    def __init__(self):
        self.window = None
        self.label = None
        self.scroll_after_id = None
        self.current_text = ""
        self.offset = 0
        self.max_display_len = 30
        self.scroll_speed = banner_scroll_speed
        self.bg_color = banner_bg_color
        self.fg_color = banner_text_color
        self.font_size = banner_font_size
        self.height = banner_height
        self.running = False
        self.hide_timer = None

    def create_window(self):
        if self.window is not None:
            try:
                self.window.destroy()
            except:
                pass
        self.window = tk.Toplevel()
        self.window.overrideredirect(True)
        self.window.attributes("-topmost", True)
        self.window.configure(bg=self.bg_color)
        screen_width = self.window.winfo_screenwidth()
        width = int(screen_width * 0.9)
        self.x = (screen_width - width) // 2
        self.y = 10
        self.window.geometry(f"{width}x{self.height}+{self.x}+{-100}")
        self.window.attributes('-alpha', max(banner_alpha, 0.01))
        self.label = tk.Label(self.window, text="", font=("微软雅黑", self.font_size, "bold"),
                              bg=self.bg_color, fg=self.fg_color, anchor="w")
        self.label.pack(fill=tk.BOTH, expand=True, padx=10, pady=5)
        self.label.bind("<Configure>", self._on_label_resize)

    def _on_label_resize(self, event):
        if not self.running:
            try:
                label_width = self.label.winfo_width()
                char_width = self.font_size * 1.2
                self.max_display_len = max(10, int(label_width / char_width))
                self.label.config(width=self.max_display_len)
            except:
                pass

    def set_scroll_speed(self, speed_ms):
        self.scroll_speed = max(30, min(500, speed_ms))
        if self.running:
            self.stop_scroll()
            self.running = True
            self._scroll_text()

    def set_alpha(self, alpha):
        if self.window:
            set_window_alpha(self.window, max(alpha, 0.01))

    def update_style(self):
        if self.window:
            self.window.configure(bg=self.bg_color)
            current_geo = self.window.geometry()
            width = int(current_geo.split('x')[0])
            self.window.geometry(f"{width}x{self.height}+{self.x}+{self.y}")
            set_window_alpha(self.window, max(banner_alpha, 0.01))
        if self.label:
            self.label.configure(bg=self.bg_color, fg=self.fg_color,
                                 font=("微软雅黑", self.font_size, "bold"))
            self._on_label_resize(None)

    def set_bg_color(self, color):
        self.bg_color = color
        self.update_style()

    def set_fg_color(self, color):
        self.fg_color = color
        self.update_style()

    def set_font_size(self, size):
        self.font_size = size
        self.update_style()

    def set_height(self, height):
        self.height = height
        self.update_style()

    def show_text(self, sender, text, on_hide_complete=None):
        full_text = f"{sender}：{text}"
        self.current_text = full_text
        if self.window is None:
            self.create_window()
        else:
            if not self.window.winfo_viewable():
                self.window.geometry(f"{self.window.winfo_width()}x{self.height}+{self.x}+{-100}")
                set_window_alpha(self.window, max(banner_alpha, 0.01))
        self.stop_scroll()
        if self.hide_timer:
            self.window.after_cancel(self.hide_timer)
            self.hide_timer = None
        self.offset = 0

        self.window.update_idletasks()
        label_width = self.label.winfo_width()
        if label_width <= 0:
            label_width = 400
        char_width = self.font_size * 1.2
        self.max_display_len = max(10, int(label_width / char_width))
        self.label.config(width=self.max_display_len)

        if not self.window.winfo_viewable():
            slide_in_banner(self.window, start_y=-100, end_y=self.y, duration=300, target_alpha=banner_alpha)

        if len(full_text) <= self.max_display_len:
            self.label.config(text=full_text)
        else:
            self.running = True
            self._scroll_text()

    def _scroll_text(self):
        if not self.running or not self.window or not self.window.winfo_viewable():
            return
        text = self.current_text
        if len(text) <= self.max_display_len:
            self.label.config(text=text)
            return
        if self.offset + self.max_display_len >= len(text):
            display = text[self.offset:] + " " * (self.offset + self.max_display_len - len(text))
        else:
            display = text[self.offset:self.offset + self.max_display_len]
        self.label.config(text=display)
        self.offset += 1
        if self.offset >= len(text):
            self.offset = 0
        self.scroll_after_id = self.window.after(self.scroll_speed, self._scroll_text)

    def stop_scroll(self):
        self.running = False
        if self.scroll_after_id:
            self.window.after_cancel(self.scroll_after_id)
            self.scroll_after_id = None

    def stop_and_hide(self, animated=True, on_complete=None):
        self.stop_scroll()
        if self.hide_timer:
            self.window.after_cancel(self.hide_timer)
            self.hide_timer = None
        if animated and self.window and self.window.winfo_viewable():
            slide_out_banner(self.window, end_y=-100, duration=300, on_complete=on_complete)
        else:
            if self.window:
                self.window.withdraw()
            if on_complete:
                on_complete()

# ================== 右下角弹窗 ==================
class PopupNotification:
    def __init__(self, root):
        self.root = root
        self.window = None
        self.hide_timer = None
        self.scroll_after_id = None
        self.bg_color = popup_bg_color
        self.text_color = popup_text_color
        self.font_size = popup_font_size
        self.max_height = popup_max_height
        self.auto_scroll = popup_auto_scroll
        self.scroll_speed = popup_vertical_scroll_speed
        self.text_widget = None
        self.current_text = ""
        self.scroll_offset = 0

    def set_bg_color(self, color):
        self.bg_color = color

    def set_text_color(self, color):
        self.text_color = color

    def set_font_size(self, size):
        self.font_size = size

    def set_max_height(self, height):
        self.max_height = max(100, height)

    def set_auto_scroll(self, enabled):
        self.auto_scroll = enabled

    def set_scroll_speed(self, speed):
        self.scroll_speed = max(20, min(500, speed))

    def show(self, title, message, auto_hide_callback=None):
        if self.window:
            self.stop_auto_scroll()
            try:
                self.window.destroy()
            except:
                pass
        self.window = tk.Toplevel(self.root)
        self.window.overrideredirect(True)
        self.window.attributes("-topmost", True)
        self.window.configure(bg=self.bg_color)

        main_frame = tk.Frame(self.window, bg=self.bg_color)
        main_frame.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)

        title_label = tk.Label(main_frame, text=title, font=("微软雅黑", self.font_size+2, "bold"),
                               bg=self.bg_color, fg=self.text_color, anchor="w")
        title_label.pack(fill=tk.X, pady=(0, 5))

        text_frame = tk.Frame(main_frame, bg=self.bg_color)
        text_frame.pack(fill=tk.BOTH, expand=True)

        self.text_widget = tk.Text(text_frame, wrap=tk.WORD, font=("微软雅黑", self.font_size),
                                   bg=self.bg_color, fg=self.text_color, relief=tk.FLAT,
                                   highlightthickness=0, borderwidth=0)
        scrollbar = tk.Scrollbar(text_frame, orient=tk.VERTICAL, command=self.text_widget.yview)
        self.text_widget.configure(yscrollcommand=scrollbar.set)

        self.text_widget.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)

        self.text_widget.insert(tk.END, message)
        self.text_widget.configure(state=tk.DISABLED)

        self.window.update_idletasks()
        text_height = self.text_widget.winfo_reqheight()
        title_height = title_label.winfo_reqheight()
        total_content_height = title_height + text_height + 20

        if total_content_height > self.max_height:
            window_height = self.max_height
            self.text_widget.configure(height=int((self.max_height - title_height - 20) / (self.font_size * 1.5)))
            if self.auto_scroll:
                self.start_auto_scroll()
        else:
            window_height = total_content_height

        lines = message.split('\n')
        max_line_len = max([len(line) for line in lines]) if lines else 10
        approx_width = min(max_line_len * self.font_size + 40, 400)
        window_width = max(approx_width, 280)

        screen_width = self.window.winfo_screenwidth()
        screen_height = self.window.winfo_screenheight()
        x = screen_width - window_width - 20
        y = screen_height - window_height - 50
        self.window.geometry(f"{window_width}x{window_height}+{x}+{y}")

        self.window.attributes('-alpha', 0.0)
        self.window.deiconify()
        fade_in(self.window, target_alpha=0.95, duration=150)

        self.auto_hide_callback = auto_hide_callback

    def start_auto_scroll(self):
        if not self.text_widget:
            return
        self.scroll_offset = 0
        self._auto_scroll_step()

    def _auto_scroll_step(self):
        if not self.window or not self.window.winfo_viewable():
            return
        if not self.auto_scroll:
            return
        total_lines = int(self.text_widget.index(tk.END).split('.')[0]) - 1
        if total_lines <= 0:
            return
        self.scroll_offset += 1
        if self.scroll_offset > total_lines:
            self.scroll_offset = 0
        self.text_widget.yview_moveto(self.scroll_offset / total_lines)
        self.scroll_after_id = self.window.after(self.scroll_speed, self._auto_scroll_step)

    def stop_auto_scroll(self):
        if self.scroll_after_id:
            self.window.after_cancel(self.scroll_after_id)
            self.scroll_after_id = None

    def hide(self, animated=True):
        self.stop_auto_scroll()
        if self.hide_timer:
            self.root.after_cancel(self.hide_timer)
            self.hide_timer = None
        if self.window:
            if animated:
                fade_out(self.window, duration=200, on_complete=lambda: self._cleanup())
            else:
                self._cleanup()

    def _cleanup(self):
        if self.window:
            self.window.destroy()
            self.window = None
        if self.auto_hide_callback:
            self.auto_hide_callback()
            self.auto_hide_callback = None

def send_notification(title, message, on_hide_complete=None):
    global popup_notification
    if popup_notification is None and root:
        popup_notification = PopupNotification(root)
    if popup_notification:
        popup_notification.show(title, message, auto_hide_callback=on_hide_complete)
        debug(f"自定义弹窗显示: {title} - {message}")

# ================== 音频系统 ==================
def stop_audio():
    global stop_flag
    stop_flag.set()
    if platform.system() == "Windows":
        try:
            ctypes.windll.winmm.mciSendStringW('stop mp3', None, 0, 0)
            ctypes.windll.winmm.mciSendStringW('close mp3', None, 0, 0)
        except:
            pass

async def edge_tts_speak(text, voice, vol):
    import edge_tts
    global stop_flag
    stop_flag.clear()
    tmp = os.path.join(tempfile.gettempdir(), "call_speech.mp3")
    for attempt in range(2):
        try:
            await asyncio.wait_for(edge_tts.Communicate(text, voice).save(tmp), timeout=15.0)
            break
        except asyncio.TimeoutError:
            debug(f"edge_tts 超时 (尝试 {attempt+1})")
            if attempt == 1:
                raise
        except Exception as e:
            debug(f"edge_tts 异常 (尝试 {attempt+1}): {e}")
            if attempt == 1:
                raise
    if platform.system() == "Windows":
        with threading.Lock():
            ctypes.windll.winmm.mciSendStringW(f'open "{tmp}" type mpegvideo alias mp3', None, 0, 0)
            ctypes.windll.winmm.mciSendStringW(f'setaudio mp3 volume to {min(max(vol, 0), 1000)}', None, 0, 0)
            ctypes.windll.winmm.mciSendStringW('play mp3', None, 0, 0)
            while True:
                if stop_flag.is_set():
                    ctypes.windll.winmm.mciSendStringW('stop mp3', None, 0, 0)
                    break
                time.sleep(0.1)
                buf = ctypes.create_unicode_buffer(128)
                ctypes.windll.winmm.mciSendStringW('status mp3 mode', buf, 128, 0)
                if buf.value == 'stopped':
                    break
            ctypes.windll.winmm.mciSendStringW('close mp3', None, 0, 0)
    elif platform.system() == "Darwin":
        subprocess.run(["afplay", tmp])
    else:
        subprocess.run(["mpg123", "-q", tmp])

def system_speak(text, voice=None):
    global stop_flag
    stop_flag.clear()
    sys_platform = platform.system()
    try:
        if sys_platform == "Windows":
            safe = text.replace('"', '`"')
            subprocess.run(["powershell", "-Command",
                            f"Add-Type -AssemblyName System.Speech; (New-Object System.Speech.Synthesis.SpeechSynthesizer).Speak('{safe}')"],
                           creationflags=subprocess.CREATE_NO_WINDOW if hasattr(subprocess, "CREATE_NO_WINDOW") else 0)
        elif sys_platform == "Darwin":
            subprocess.run(["say", text] + (["-v", voice] if voice else []))
        else:
            subprocess.run(["espeak", text] + (["-v", voice] if voice else []))
    except Exception as e:
        debug(f"系统语音异常: {e}")

def speak(text, voice="zh-CN-XiaoxiaoNeural", vol=1000):
    if not speech_enabled:
        debug(f"语音播报已禁用，跳过: {text}")
        return
    if EDGE_AVAILABLE:
        try:
            loop = asyncio.new_event_loop()
            asyncio.set_event_loop(loop)
            loop.run_until_complete(edge_tts_speak(text, voice, vol))
            loop.close()
            debug(f"edge_tts 播报成功: {text}")
            return
        except Exception as e:
            debug(f"edge_tts 播报失败: {e}，回退到系统语音")
    system_speak(text, voice)

voice_options = [
    {"id": "zh-CN-XiaoxiaoNeural", "name": "中文女声 (晓晓)"},
    {"id": "zh-CN-YunxiNeural", "name": "中文男声 (云希)"},
    {"id": "en-US-JennyNeural", "name": "英文女声 (Jenny)"},
    {"id": "en-US-GuyNeural", "name": "英文男声 (Guy)"},
]

# ================== 设备码 ==================
def get_device_code():
    path = os.path.join(BASE_DIR, "device_code.txt")
    try:
        if os.path.exists(path):
            with open(path, "r") as f:
                return f.read().strip()
    except:
        pass
    code = uuid.uuid4().hex[:8].upper()
    try:
        with open(path, "w") as f:
            f.write(code)
    except:
        pass
    return code

DEVICE_CODE = get_device_code()

# ================== 队列管理器 ==================
class TaskQueue:
    def __init__(self):
        self.tasks = []
        self.lock = threading.Lock()
        self.cond = threading.Condition(self.lock)
        self.running = True
        self.paused = False
        self.mode = "manual"
        self.interval = 2
        self.manual_event = threading.Event()
        self.worker = threading.Thread(target=self._worker, daemon=True)
        self.worker.start()

    def add(self, task):
        with self.cond:
            self.tasks.append(task)
            self.cond.notify()

    def clear(self):
        with self.cond:
            self.tasks.clear()
            self.paused = False
            self.cond.notify_all()
            stop_audio()

    def remove(self, idx):
        with self.cond:
            if 0 <= idx < len(self.tasks):
                del self.tasks[idx]

    def move(self, idx, d):
        with self.cond:
            nxt = idx + d
            if 0 <= idx < len(self.tasks) and 0 <= nxt < len(self.tasks):
                self.tasks[idx], self.tasks[nxt] = self.tasks[nxt], self.tasks[idx]

    def set_mode(self, mode, interval=None):
        with self.cond:
            self.mode = mode
            if interval is not None:
                self.interval = max(1, interval)
            self.cond.notify_all()

    def manual_next(self):
        self.manual_event.set()

    def pause(self):
        with self.cond:
            self.paused = True
            self.cond.notify_all()

    def resume(self):
        with self.cond:
            self.paused = False
            self.cond.notify_all()

    def get_status(self):
        with self.lock:
            return {"tasks": self.tasks, "mode": self.mode, "interval": self.interval, "queue_count": len(self.tasks)}

    def _update_status(self):
        if root and classic_window:
            s = self.get_status()
            try:
                root.after(0, classic_window.update_status, s["queue_count"],
                           "播放中" if (self.running and not self.paused) else "暂停" if self.paused else "空闲",
                           "手动" if self.mode == "manual" else "自动")
            except:
                pass

    def _worker(self):
        while self.running:
            task = None
            with self.cond:
                if self.mode == "auto":
                    while self.running and (not self.tasks or self.paused):
                        self.cond.wait()
                    if not self.running:
                        break
                    task = self.tasks.pop(0)
                else:
                    while self.running and self.mode == "manual":
                        if self.tasks and not self.paused:
                            self.cond.release()
                            self.manual_event.wait()
                            self.manual_event.clear()
                            self.cond.acquire()
                            if not self.running or self.mode != "manual" or not self.tasks or self.paused:
                                continue
                            task = self.tasks.pop(0)
                            break
                        else:
                            self.cond.wait()
            if task:
                self._execute_task(task)
                if self.mode == "auto" and self.tasks and self.running and not self.paused:
                    time.sleep(self.interval)
            self._update_status()

    def _execute_task(self, task):
        global display_mode, speech_enabled, banner_window
        text = task["text"]
        sender = sender_name
        always = task.get("always_show", False)

        if display_mode == "classic":
            if classic_window:
                classic_window.show_text(text)
            if speech_enabled:
                self._play_sound(task)
            if not always:
                if root and classic_window:
                    root.after(1000, lambda: classic_window.schedule_hide(500, animated=True))
        elif display_mode == "banner":
            if banner_window:
                banner_window.show_text(sender, text)
            if speech_enabled:
                self._play_sound(task)
                if banner_window and banner_window.window:
                    root.after(1000, lambda: banner_window.stop_and_hide(animated=True))
            else:
                if banner_window and banner_window.window:
                    root.after(10000, lambda: banner_window.stop_and_hide(animated=True))
        elif display_mode == "notification":
            send_notification(title=f"{sender}", message=text, on_hide_complete=None)
            if speech_enabled:
                self._play_sound(task)
                if popup_notification:
                    root.after(0, lambda: popup_notification.hide(animated=True))
            else:
                if popup_notification:
                    root.after(5000, lambda: popup_notification.hide(animated=True))

    def _play_sound(self, task):
        t = task["type"]
        if t == "normal":
            speak(task["text"], task.get("voice", "zh-CN-XiaoxiaoNeural"), task.get("volume", 1000))
        elif t == "repeat":
            for _ in range(task.get("repeat_count", 1)):
                if self.paused or not self.running:
                    break
                speak(task["text"], task.get("voice"), task.get("volume", 1000))
                if task.get("interval", 0) > 0:
                    time.sleep(task["interval"])
        elif t == "timed":
            trigger = task.get("trigger_time")
            if trigger and trigger > datetime.datetime.now():
                wait = (trigger - datetime.datetime.now()).total_seconds()
                end = time.time() + wait
                while time.time() < end and self.running and not self.paused:
                    time.sleep(0.5)
            if not self.running or self.paused:
                return
            speak(task["text"], task.get("voice"), task.get("volume", 1000))

# ================== Flask 应用 ==================
app = Flask(__name__)
app.secret_key = os.urandom(24)

def login_required(f):
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if not session.get('logged_in'):
            return redirect(url_for('login_page'))
        return f(*args, **kwargs)
    return decorated_function

def api_login_required(f):
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if not session.get('logged_in'):
            return jsonify({'error': '未授权，请先登录'}), 401
        return f(*args, **kwargs)
    return decorated_function

def record_user_action():
    global last_user_action_time
    last_user_action_time = datetime.datetime.now()
    debug(f"用户操作记录: {last_user_action_time}")

# 登录页面
LOGIN_TEMPLATE = """
<!DOCTYPE html>
<html>
<head>
    <meta charset="utf-8">
    <title>班呼ClassCall - 登录</title>
    <style>
        * { margin:0; padding:0; box-sizing:border-box; }
        body { 
            font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, '微软雅黑', sans-serif;
            background: #f5f7fa;
            height: 100vh;
            display: flex;
            justify-content: center;
            align-items: center;
        }
        .login-card {
            background: white;
            border-radius: 24px;
            padding: 40px;
            width: 350px;
            box-shadow: 0 8px 20px rgba(0,0,0,0.08);
            text-align: center;
        }
        .login-card h1 {
            color: #2d3748;
            margin-bottom: 8px;
            font-size: 28px;
        }
        .login-card p {
            color: #718096;
            margin-bottom: 30px;
            font-size: 14px;
        }
        input {
            width: 100%;
            padding: 12px 16px;
            margin: 12px 0;
            border: 1px solid #e2e8f0;
            border-radius: 12px;
            font-size: 16px;
            transition: 0.2s;
        }
        input:focus {
            outline: none;
            border-color: #667eea;
            box-shadow: 0 0 0 3px rgba(102,126,234,0.2);
        }
        button {
            width: 100%;
            padding: 12px;
            background: #667eea;
            color: white;
            border: none;
            border-radius: 12px;
            font-size: 16px;
            font-weight: 600;
            cursor: pointer;
            margin-top: 12px;
            transition: 0.2s;
        }
        button:hover {
            background: #5a67d8;
        }
        .error {
            color: #e53e3e;
            margin-top: 12px;
            font-size: 14px;
        }
        .device-code {
            margin-top: 20px;
            font-size: 12px;
            color: #a0aec0;
        }
    </style>
</head>
<body>
    <div class="login-card">
        <h1>📢 班呼ClassCall</h1>
        <p>请输入密码以访问控制面板</p>
        <form method="POST" action="/login">
            <input type="password" name="password" placeholder="密码" autofocus>
            <button type="submit">登录</button>
            {% if error %}
            <div class="error">{{ error }}</div>
            {% endif %}
        </form>
        <div class="device-code">设备码: {{ device_code }}</div>
    </div>
</body>
</html>
"""

@app.route('/login', methods=['GET', 'POST'])
def login_page():
    error = None
    if request.method == 'POST':
        pwd = request.form.get('password', '')
        if pwd == WEB_PASSWORD:
            session['logged_in'] = True
            return redirect(url_for('dashboard'))
        else:
            error = "密码错误，请重试"
    return render_template_string(LOGIN_TEMPLATE, device_code=DEVICE_CODE, error=error)

@app.route('/logout')
def logout():
    session.pop('logged_in', None)
    return redirect(url_for('login_page'))

def parse_roster_file(file):
    filename = file.filename
    ext = os.path.splitext(filename)[1].lower()
    names = []
    if ext == '.txt':
        content = file.read().decode('utf-8', errors='ignore')
        for line in content.splitlines():
            line = line.strip()
            if line:
                names.append(line)
    elif ext == '.csv':
        content = file.read().decode('utf-8', errors='ignore')
        import csv
        reader = csv.reader(content.splitlines())
        for row in reader:
            if row:
                names.append(row[0].strip())
    elif ext in ('.xlsx', '.xls'):
        if not OPENPYXL_AVAILABLE:
            raise Exception("未安装 openpyxl，无法读取 Excel 文件")
        import openpyxl
        wb = openpyxl.load_workbook(file)
        sheet = wb.active
        for row in sheet.iter_rows(values_only=True):
            if row and row[0]:
                names.append(str(row[0]).strip())
    else:
        raise Exception("不支持的文件格式，请上传 .txt, .csv, .xlsx 文件")
    seen = set()
    unique_names = []
    for n in names:
        if n not in seen:
            seen.add(n)
            unique_names.append(n)
    return unique_names

@app.route('/upload_roster', methods=['POST'])
@api_login_required
def upload_roster():
    if 'file' not in request.files:
        return jsonify({'error': '未选择文件'}), 400
    file = request.files['file']
    if file.filename == '':
        return jsonify({'error': '文件名为空'}), 400
    try:
        names = parse_roster_file(file)
        return jsonify({'names': names})
    except Exception as e:
        debug(f"解析名单失败: {e}")
        return jsonify({'error': str(e)}), 400

@app.route('/change_web_password', methods=['POST'])
@api_login_required
def change_web_password():
    global WEB_PASSWORD, config
    data = request.get_json()
    old = data.get('old_password', '')
    new = data.get('new_password', '')
    if not old or not new:
        return jsonify({'error': '旧密码和新密码不能为空'}), 400
    if old != WEB_PASSWORD:
        return jsonify({'error': '旧密码错误'}), 400
    if len(new) < 4:
        return jsonify({'error': '新密码长度至少4位'}), 400
    WEB_PASSWORD = new
    config['web_password'] = new
    save_config(config)
    record_user_action()
    return jsonify({'status': 'ok', 'message': '网页密码修改成功'})

@app.route('/change_exit_password', methods=['POST'])
@api_login_required
def change_exit_password():
    global EXIT_PASSWORD, config
    data = request.get_json()
    old = data.get('old_password', '')
    new = data.get('new_password', '')
    if not old or not new:
        return jsonify({'error': '旧密码和新密码不能为空'}), 400
    if old != EXIT_PASSWORD:
        return jsonify({'error': '旧密码错误'}), 400
    if len(new) < 4:
        return jsonify({'error': '新密码长度至少4位'}), 400
    EXIT_PASSWORD = new
    config['exit_password'] = new
    save_config(config)
    record_user_action()
    return jsonify({'status': 'ok', 'message': '退出密码修改成功'})

# ================== 完整功能HTML模板（默认展开高级设置） ==================
HTML_TEMPLATE = r"""<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>班呼ClassCall - 控制面板</title>
<style>
    * { margin:0; padding:0; box-sizing:border-box; }
    body { font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', '微软雅黑', sans-serif; background: #f5f7fa; padding: 20px; color: #2d3748; }
    .container { max-width: 1200px; margin: 0 auto; }
    .card { background: white; border-radius: 24px; padding: 20px; margin-bottom: 20px; box-shadow: 0 2px 8px rgba(0,0,0,0.04); border: 1px solid #edf2f7; }
    .header { display: flex; justify-content: space-between; align-items: center; margin-bottom: 16px; flex-wrap: wrap; gap: 10px; }
    .device-badge { background: #e2e8f0; color: #4a5568; padding: 6px 16px; border-radius: 30px; font-size: 13px; font-weight: 600; }
    h2 { font-size: 20px; margin-bottom: 12px; color: #2d3748; border-left: 4px solid #667eea; padding-left: 12px; }
    textarea { width: 100%; height: 100px; border: 1px solid #e2e8f0; border-radius: 16px; padding: 12px; font-size: 16px; resize: vertical; background: white; }
    textarea:focus { outline: none; border-color: #667eea; box-shadow: 0 0 0 3px rgba(102,126,234,0.2); }
    .row { display: flex; align-items: center; gap: 12px; margin: 14px 0; flex-wrap: wrap; }
    select, input[type=range], input[type=number], input[type=time] {
        flex:1; padding: 10px 14px; border: 1px solid #e2e8f0; border-radius: 14px; font-size: 15px; background: white; min-width: 0;
    }
    .btn {
        border: none;
        border-radius: 40px;
        padding: 10px 20px;
        font-size: 14px;
        cursor: pointer;
        font-weight: 600;
        transition: 0.2s;
        background: #edf2f7;
        color: #4a5568;
    }
    .btn-primary { background: #667eea; color: white; }
    .btn-primary:hover { background: #5a67d8; transform: translateY(-1px); }
    .btn-danger { background: #fc8181; color: white; }
    .btn-danger:hover { background: #f56565; }
    .btn-row { display: flex; gap: 10px; margin: 10px 0; flex-wrap: wrap; }
    .queue-item { display: flex; align-items: center; gap: 8px; padding: 10px; background: #f8fafc; border-radius: 16px; margin: 6px 0; border-left: 4px solid #667eea; }
    .name-list { display: flex; flex-wrap: wrap; gap: 8px; margin-top: 12px; max-height: 300px; overflow-y: auto; }
    .name-item { background: #e9ecef; padding: 8px 16px; border-radius: 30px; cursor: pointer; transition: 0.1s; font-size: 14px; }
    .name-item:hover { background: #cbd5e1; transform: scale(1.02); }
    .search-box { width: 100%; padding: 10px; margin-bottom: 12px; border: 1px solid #e2e8f0; border-radius: 30px; }
    .modal { display:none; position:fixed; top:0; left:0; width:100%; height:100%; background:rgba(0,0,0,0.5); z-index:1000; justify-content:center; align-items:center; }
    .modal-content { background:white; border-radius: 32px; padding: 24px; width:90%; max-width: 500px; max-height:80vh; overflow:auto; }
    .setting-group { transition: all 0.3s; margin-top: 12px; }
    .color-preview { display: inline-block; width: 32px; height: 32px; border-radius: 16px; border: 2px solid white; box-shadow: 0 1px 3px rgba(0,0,0,0.2); margin-left: 8px; vertical-align: middle; }
    hr { margin: 16px 0; border: none; height: 1px; background: #e2e8f0; }
    .collapse-toggle { color: #667eea; font-weight: 600; text-align: center; cursor: pointer; margin: 12px 0; background: #f1f5f9; padding: 8px; border-radius: 30px; }
    @media (max-width: 640px) { .btn-row button { flex: 1; } }
</style>
</head>
<body>
<div class="container">
<div class="card">
    <div class="header">
        <span style="font-size:24px;font-weight:800;">📢 班呼ClassCall</span>
        <span class="device-badge">设备码 {{ device_code }}</span>
    </div>
    <div class="row">
        <label>发件人:</label><input type="text" id="senderName" value="{{ sender_name }}" style="flex:2;">
        <button class="btn btn-primary" onclick="setSender()">设置</button>
        <button class="btn btn-primary" onclick="openRosterModal()">📋 导入名单</button>
    </div>
    <textarea id="text" placeholder="输入叫号内容..."></textarea>
    <div class="row"><select id="voice">{% for v in voices %}<option value="{{ v.id }}">{{ v.name }}</option>{% endfor %}</select></div>
    <div class="row"><span>🔊 音量</span><input type="range" id="vol" min="0" max="500" value="100" step="10"><span id="volVal">100%</span></div>
    <div class="row">
        <label><input type="checkbox" id="alwaysShow"> 常显</label>
        <input type="number" id="showDuration" value="0" min="0" step="1" style="width:80px;" placeholder="秒">
        <label>显示秒数(0=自动)</label>
    </div>
    <div class="btn-row">
        <button class="btn btn-primary" onclick="addToQueue()">📋 加入队列</button>
        <button class="btn btn-primary" onclick="startRepeat()">🔄 重复朗读</button>
    </div>
    <div class="btn-row">
        <button class="btn btn-primary" onclick="addTimed()">⏰ 定时朗读</button>
        <button class="btn btn-danger" onclick="stopAll()">⏹️ 停止</button>
    </div>
</div>

<div class="card">
    <div style="font-weight:700;">📋 排队控制 <span id="qc">(0)</span></div>
    <div class="row"><select id="queueMode" onchange="setMode()"><option value="manual" selected>手动</option><option value="auto">自动</option></select></div>
    <div class="row" id="autoSet" style="display:none;"><span>间隔(秒)</span><input type="number" id="autoInterval" value="2" min="1"></div>
    <div class="btn-row">
        <button id="nextBtn" class="btn btn-primary" onclick="manualNext()">🔽 下一条</button>
        <button id="autoStartBtn" class="btn btn-primary" style="display:none;" onclick="startAuto()">▶️ 开始自动</button>
        <button class="btn" onclick="pauseQueue()">⏸️ 暂停</button>
        <button class="btn btn-danger" onclick="clearQueue()">🗑️ 清空</button>
    </div>
    <div id="queueList"></div>
</div>

<!-- 高级设置区域默认展开（去掉折叠，直接显示） -->
<div class="card">
    <h2>⚙️ 高级设置 & 密码修改</h2>
    <div class="row"><label>显示模式：</label><select id="displayMode" onchange="onDisplayModeChange()"><option value="classic" {{ 'selected' if display_mode=='classic' else '' }}>传统窗口</option><option value="banner" {{ 'selected' if display_mode=='banner' else '' }}>顶部横幅（滚动）</option><option value="notification" {{ 'selected' if display_mode=='notification' else '' }}>右下角弹窗</option></select><button class="btn" onclick="setDisplayMode()">应用</button></div>
    <div class="row"><label><input type="checkbox" id="speechEnabled" {{ 'checked' if speech_enabled else '' }}> 启用语音播报</label><button class="btn" onclick="setSpeechEnabled()">保存</button></div>
    
    <div id="classicSettings" class="setting-group" style="display: {{ 'block' if display_mode=='classic' else 'none' }};"><hr>
        <div class="btn-row"><button class="btn" onclick="toggleGreen()">🟢 绿幕</button><button class="btn" onclick="toggleFullscreen()">🖥️ 全屏</button></div>
        <div class="row"><button class="btn" onclick="openColorPicker('bg')">背景颜色</button><span id="bgColorPreview" class="color-preview" style="background-color:{{ bg_color }}"></span></div>
        <div class="row"><button class="btn" onclick="openColorPicker('text')">文字颜色</button><span id="textColorPreview" class="color-preview" style="background-color:{{ classic_text_color }}"></span></div>
        <div class="row"><label>窗口透明度</label><input type="range" id="classicAlpha" min="0" max="100" value="{{ classic_alpha*100|int }}"><span id="classicAlphaVal">{{ "%.0f"|format(classic_alpha*100) }}%</span><button class="btn" onclick="setClassicAlpha()">应用</button></div>
        <div class="row"><select id="fontFamily"><option>微软雅黑</option><option>黑体</option><option>宋体</option><option>楷体</option><option>Arial</option></select></div>
        <div class="row"><input type="number" id="fontSize" value="48" min="20" max="150"></div>
        <button class="btn" onclick="applyFontSettings()">应用字体设置</button>
    </div>

    <div id="bannerSettings" class="setting-group" style="display: {{ 'block' if display_mode=='banner' else 'none' }};"><hr>
        <div class="row"><button class="btn" onclick="openColorPicker('bannerBg')">横幅背景色</button><span id="bannerBgPreview" class="color-preview" style="background-color:{{ banner_bg_color }}"></span></div>
        <div class="row"><button class="btn" onclick="openColorPicker('bannerText')">横幅文字颜色</button><span id="bannerTextPreview" class="color-preview" style="background-color:{{ banner_text_color }}"></span></div>
        <div class="row"><label>横幅高度</label><input type="number" id="bannerHeight" value="{{ banner_height }}" min="40" max="200"><label>字体大小</label><input type="number" id="bannerFontSize" value="{{ banner_font_size }}" min="12" max="48"></div>
        <div class="row"><label>滚动速度(毫秒)</label><input type="range" id="bannerScrollSpeed" min="30" max="500" value="{{ banner_scroll_speed }}" step="10"><span id="bannerScrollSpeedVal">{{ banner_scroll_speed }}ms</span><button class="btn" onclick="setBannerScrollSpeed()">应用</button></div>
        <div class="row"><label>横幅透明度</label><input type="range" id="bannerAlpha" min="0" max="100" value="{{ banner_alpha*100|int }}"><span id="bannerAlphaVal">{{ "%.0f"|format(banner_alpha*100) }}%</span><button class="btn" onclick="setBannerAlpha()">应用</button></div>
        <div class="btn-row"><button class="btn" onclick="setBannerStyle()">应用尺寸</button></div>
    </div>

    <div id="popupSettings" class="setting-group" style="display: {{ 'block' if display_mode=='notification' else 'none' }};"><hr>
        <div class="row"><button class="btn" onclick="openColorPicker('popupBg')">弹窗背景色</button><span id="popupBgPreview" class="color-preview" style="background-color:{{ popup_bg_color }}"></span></div>
        <div class="row"><button class="btn" onclick="openColorPicker('popupText')">弹窗文字颜色</button><span id="popupTextPreview" class="color-preview" style="background-color:{{ popup_text_color }}"></span></div>
        <div class="row"><label>弹窗字体大小</label><input type="number" id="popupFontSize" value="{{ popup_font_size }}" min="10" max="30"><button class="btn" onclick="setPopupStyle()">应用</button></div>
        <div class="row"><label>弹窗最大高度(px)</label><input type="number" id="popupMaxHeight" value="{{ popup_max_height }}" min="100" max="600"><button class="btn" onclick="setPopupMaxHeight()">应用</button></div>
        <div class="row"><label><input type="checkbox" id="popupAutoScroll" {{ 'checked' if popup_auto_scroll else '' }}> 启用自动竖向滚动</label><button class="btn" onclick="setPopupAutoScroll()">保存</button></div>
        <div class="row"><label>竖向滚动速度(ms/行)</label><input type="range" id="popupVerticalScrollSpeed" min="20" max="500" value="{{ popup_vertical_scroll_speed }}" step="10"><span id="popupVerticalScrollSpeedVal">{{ popup_vertical_scroll_speed }}ms</span><button class="btn" onclick="setPopupScrollSpeed()">应用</button></div>
    </div>
    <hr>
    <div class="row"><button class="btn btn-primary" onclick="openChangePasswordModal()">🔐 修改密码</button></div>
</div>
</div>

<!-- 模态框 -->
<div id="rosterModal" class="modal"><div class="modal-content"><h3>导入班级名单</h3><div class="row"><input type="file" id="rosterFile" accept=".txt,.csv,.xlsx"><button class="btn btn-primary" onclick="uploadRoster()">上传并解析</button></div><div id="rosterStatus"></div><input type="text" id="searchName" class="search-box" placeholder="🔍 搜索名字..." onkeyup="filterNames()"><div id="nameListContainer" class="name-list"></div><div class="btn-row"><button class="btn" onclick="closeRosterModal()">关闭</button></div></div></div>
<div id="colorModal" class="modal"><div class="modal-content"><h3 id="colorModalTitle">选择颜色</h3><input type="color" id="colorPicker" value="#000000"><div class="btn-row"><button class="btn btn-primary" id="confirmColorBtn">确定</button><button class="btn" onclick="closeColorModal()">取消</button></div></div></div>
<div id="repeatModal" class="modal"><div class="modal-content"><h3>🔄 重复朗读</h3><div class="row"><label>次数</label><input id="repeatCnt" value="3" min="1"></div><div class="row"><label>间隔</label><input id="repeatInterval" value="2" min="0"></div><div class="btn-row"><button class="btn btn-primary" onclick="confirmRepeat()">开始</button><button class="btn" onclick="closeModal('repeatModal')">取消</button></div></div></div>
<div id="timedModal" class="modal"><div class="modal-content"><h3>⏰ 定时朗读</h3><textarea id="timedText" placeholder="内容..."></textarea><div class="row"><select id="timedType" onchange="toggleTimed()"><option value="time">指定时间</option><option value="countdown">倒计时(秒)</option></select></div><div id="timeInput"><input type="time" id="timedTime"></div><div id="cntInput" style="display:none;"><input type="number" id="cntSec" value="10" min="1"></div><div class="btn-row"><button class="btn btn-primary" onclick="confirmTimed()">添加</button><button class="btn" onclick="closeModal('timedModal')">取消</button></div></div></div>
<div id="changePwdModal" class="modal"><div class="modal-content"><h3>🔐 修改密码</h3><div class="row"><label>密码类型:</label><select id="pwdType"><option value="web">网页登录密码</option><option value="exit">程序退出密码</option></select></div><div class="row"><label>旧密码:</label><input type="password" id="oldPwd"></div><div class="row"><label>新密码:</label><input type="password" id="newPwd"></div><div class="row"><label>确认新密码:</label><input type="password" id="confirmPwd"></div><div class="btn-row"><button class="btn btn-primary" onclick="changePassword()">确认修改</button><button class="btn" onclick="closeModal('changePwdModal')">取消</button></div></div></div>

<script>
const DEVICE = "{{ device_code }}";
let currentColorTarget = null;
let allNames = [], currentNames = [];

// 保存原始fetch
const originalFetch = window.fetch;
window.fetch = function(...args) {
    return originalFetch.apply(this, args).then(response => {
        if (response.status === 401) {
            window.location.href = '/logout';
            return Promise.reject('Unauthorized');
        }
        return response;
    });
};

document.getElementById('vol').oninput = e => document.getElementById('volVal').innerText = e.target.value + '%';
function getVol() { return Math.min(Math.floor(parseInt(document.getElementById('vol').value)*1000/500), 1000); }

function rgbToHex(rgb) {
    if (!rgb) return '#000000';
    let result = rgb.match(/\d+/g);
    if (!result) return '#000000';
    return '#' + result.slice(0,3).map(x => parseInt(x).toString(16).padStart(2,'0')).join('');
}

function openColorPicker(target) {
    currentColorTarget = target;
    let modal = document.getElementById('colorModal');
    let title = '';
    let defaultColor = '#000000';
    switch(target) {
        case 'bg': title = '传统窗口背景色'; defaultColor = document.getElementById('bgColorPreview').style.backgroundColor; break;
        case 'text': title = '传统窗口文字颜色'; defaultColor = document.getElementById('textColorPreview').style.backgroundColor; break;
        case 'bannerBg': title = '横幅背景色'; defaultColor = document.getElementById('bannerBgPreview').style.backgroundColor; break;
        case 'bannerText': title = '横幅文字颜色'; defaultColor = document.getElementById('bannerTextPreview').style.backgroundColor; break;
        case 'popupBg': title = '弹窗背景色'; defaultColor = document.getElementById('popupBgPreview').style.backgroundColor; break;
        case 'popupText': title = '弹窗文字颜色'; defaultColor = document.getElementById('popupTextPreview').style.backgroundColor; break;
    }
    document.getElementById('colorModalTitle').innerText = title;
    let colorInput = document.getElementById('colorPicker');
    if (defaultColor && defaultColor !== '') colorInput.value = rgbToHex(defaultColor);
    else colorInput.value = '#000000';
    modal.style.display = 'flex';
}
function closeColorModal() { document.getElementById('colorModal').style.display = 'none'; currentColorTarget = null; }
document.getElementById('confirmColorBtn').onclick = async function() {
    let color = document.getElementById('colorPicker').value;
    if (!currentColorTarget) return;
    let url = '', body = {};
    switch(currentColorTarget) {
        case 'bg': url = '/set_bg_color'; body = {color:color}; break;
        case 'text': url = '/set_font'; let family = document.getElementById('fontFamily').value; let size = parseInt(document.getElementById('fontSize').value); body = {family, size, color}; break;
        case 'bannerBg': url = '/set_banner_colors'; let fg = document.getElementById('bannerTextPreview').style.backgroundColor || '#ffffff'; body = {bg:color, fg:fg}; break;
        case 'bannerText': url = '/set_banner_colors'; let bg2 = document.getElementById('bannerBgPreview').style.backgroundColor || '#2c3e50'; body = {bg:bg2, fg:color}; break;
        case 'popupBg': url = '/set_popup_bg_color'; body = {color:color}; break;
        case 'popupText': url = '/set_popup_text_color'; body = {color:color}; break;
    }
    if(url) await fetch(url, {method:'POST', headers:{'Content-Type':'application/json'}, body:JSON.stringify(body)});
    if(currentColorTarget === 'bg') document.getElementById('bgColorPreview').style.backgroundColor = color;
    else if(currentColorTarget === 'text') document.getElementById('textColorPreview').style.backgroundColor = color;
    else if(currentColorTarget === 'bannerBg') document.getElementById('bannerBgPreview').style.backgroundColor = color;
    else if(currentColorTarget === 'bannerText') document.getElementById('bannerTextPreview').style.backgroundColor = color;
    else if(currentColorTarget === 'popupBg') document.getElementById('popupBgPreview').style.backgroundColor = color;
    else if(currentColorTarget === 'popupText') document.getElementById('popupTextPreview').style.backgroundColor = color;
    closeColorModal();
    alert('颜色已更新');
};

async function setSender() { let n=document.getElementById('senderName').value.trim(); if(!n)return; await fetch('/set_sender',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({name:n})}); alert('已设置'); }
async function refreshQueue() { let r=await fetch('/queue_status'); let d=await r.json(); document.getElementById('qc').innerText=`(${d.queue_count})`; let h=d.tasks.map((t,i)=>`<div class="queue-item"><span style="flex:1">${i+1}. ${t.text}</span><button class="btn" onclick="moveTask(${i},-1)">↑</button><button class="btn" onclick="moveTask(${i},1)">↓</button><button class="btn btn-danger" onclick="removeTask(${i})">✕</button></div>`).join(''); document.getElementById('queueList').innerHTML=h; if(d.mode==='auto'){ document.getElementById('autoSet').style.display='flex'; document.getElementById('nextBtn').style.display='none'; document.getElementById('autoStartBtn').style.display='inline-block'; }else{ document.getElementById('autoSet').style.display='none'; document.getElementById('nextBtn').style.display='inline-block'; document.getElementById('autoStartBtn').style.display='none'; } }
function getTaskOpts() { return { always_show: document.getElementById('alwaysShow').checked, show_duration: parseInt(document.getElementById('showDuration').value)||0 }; }
async function addToQueue() { let t=document.getElementById('text').value.trim(); if(!t)return; await fetch('/add_task',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({type:'normal',text:t,voice:document.getElementById('voice').value,volume:getVol(),...getTaskOpts()})}); refreshQueue(); }
function startRepeat() { document.getElementById('repeatModal').style.display='flex'; }
async function confirmRepeat() { let t=document.getElementById('text').value.trim(); if(!t)return; await fetch('/add_task',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({type:'repeat',text:t,voice:document.getElementById('voice').value,volume:getVol(),repeat_count:parseInt(document.getElementById('repeatCnt').value)||1,interval:parseInt(document.getElementById('repeatInterval').value)||0,...getTaskOpts()})}); closeModal('repeatModal'); refreshQueue(); }
function addTimed() { document.getElementById('timedModal').style.display='flex'; }
function toggleTimed() { document.getElementById('timeInput').style.display=document.getElementById('timedType').value==='time'?'block':'none'; document.getElementById('cntInput').style.display=document.getElementById('timedType').value==='countdown'?'block':'none'; }
async function confirmTimed() { let t=document.getElementById('timedText').value.trim(); if(!t)return; let type=document.getElementById('timedType').value, trigger=null; if(type==='time'){ let v=document.getElementById('timedTime').value; if(!v)return; let [h,m]=v.split(':'); let d=new Date(); trigger=new Date(d.getFullYear(),d.getMonth(),d.getDate(),h,m,0); if(trigger<=d) trigger.setDate(trigger.getDate()+1); }else{ let s=parseInt(document.getElementById('cntSec').value)||10; trigger=new Date(Date.now()+s*1000); } await fetch('/add_task',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({type:'timed',text:t,voice:document.getElementById('voice').value,volume:getVol(),trigger_time:trigger.toISOString(),...getTaskOpts()})}); closeModal('timedModal'); refreshQueue(); }
async function stopAll() { await fetch('/stop',{method:'POST'}); refreshQueue(); }
async function setMode() { let m=document.getElementById('queueMode').value; await fetch('/set_mode',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({mode:m,interval:parseInt(document.getElementById('autoInterval').value)||2})}); refreshQueue(); }
async function manualNext() { await fetch('/manual_next',{method:'POST'}); refreshQueue(); }
async function startAuto() { await fetch('/set_mode',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({mode:'auto',interval:parseInt(document.getElementById('autoInterval').value)||2})}); refreshQueue(); }
async function pauseQueue() { await fetch('/pause_queue',{method:'POST'}); refreshQueue(); }
async function clearQueue() { await fetch('/clear_queue',{method:'POST'}); refreshQueue(); }
async function removeTask(i) { await fetch('/remove_task',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({index:i})}); refreshQueue(); }
async function moveTask(i,d) { await fetch('/move_task',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({index:i,direction:d})}); refreshQueue(); }
function closeModal(id) { document.getElementById(id).style.display='none'; }
function onDisplayModeChange() { let mode = document.getElementById('displayMode').value; document.getElementById('classicSettings').style.display = mode === 'classic' ? 'block' : 'none'; document.getElementById('bannerSettings').style.display = mode === 'banner' ? 'block' : 'none'; document.getElementById('popupSettings').style.display = mode === 'notification' ? 'block' : 'none'; }
async function setDisplayMode() { let mode = document.getElementById('displayMode').value; await fetch('/set_display_mode', {method:'POST', headers:{'Content-Type':'application/json'}, body:JSON.stringify({mode:mode})}); alert('显示模式已切换，下次叫号生效'); onDisplayModeChange(); }
async function setSpeechEnabled() { let enabled = document.getElementById('speechEnabled').checked; await fetch('/set_speech_enabled', {method:'POST', headers:{'Content-Type':'application/json'}, body:JSON.stringify({enabled:enabled})}); alert('语音开关已保存'); }
async function toggleGreen() { await fetch('/toggle_green',{method:'POST'}); }
async function toggleFullscreen() { await fetch('/toggle_fullscreen',{method:'POST'}); }
async function setClassicAlpha() { let val = parseInt(document.getElementById('classicAlpha').value) / 100; await fetch('/set_classic_alpha', {method:'POST', headers:{'Content-Type':'application/json'}, body:JSON.stringify({alpha:val})}); alert('透明度已更新'); }
document.getElementById('classicAlpha').oninput = function() { document.getElementById('classicAlphaVal').innerText = this.value + '%'; }
async function setBannerAlpha() { let val = parseInt(document.getElementById('bannerAlpha').value) / 100; await fetch('/set_banner_alpha', {method:'POST', headers:{'Content-Type':'application/json'}, body:JSON.stringify({alpha:val})}); alert('横幅透明度已更新'); }
document.getElementById('bannerAlpha').oninput = function() { document.getElementById('bannerAlphaVal').innerText = this.value + '%'; }
async function setBannerScrollSpeed() { let speed = parseInt(document.getElementById('bannerScrollSpeed').value); await fetch('/set_banner_scroll_speed', {method:'POST', headers:{'Content-Type':'application/json'}, body:JSON.stringify({speed:speed})}); document.getElementById('bannerScrollSpeedVal').innerText = speed + 'ms'; alert('滚动速度已更新'); }
document.getElementById('bannerScrollSpeed').oninput = function() { document.getElementById('bannerScrollSpeedVal').innerText = this.value + 'ms'; }
async function setBannerStyle() { let height = parseInt(document.getElementById('bannerHeight').value); let fontSize = parseInt(document.getElementById('bannerFontSize').value); await fetch('/set_banner_style', {method:'POST', headers:{'Content-Type':'application/json'}, body:JSON.stringify({height:height, font_size:fontSize})}); alert('横幅尺寸已更新'); }
async function applyFontSettings() { let family = document.getElementById('fontFamily').value; let size = parseInt(document.getElementById('fontSize').value); let color = document.getElementById('textColorPreview').style.backgroundColor; if (color && color !== '') color = rgbToHex(color); else color = '#ffffff'; await fetch('/set_font', {method:'POST', headers:{'Content-Type':'application/json'}, body:JSON.stringify({family:family, size:size, color:color})}); alert('字体设置已应用'); }
async function setPopupStyle() { let fontSize = parseInt(document.getElementById('popupFontSize').value); await fetch('/set_popup_style', {method:'POST', headers:{'Content-Type':'application/json'}, body:JSON.stringify({font_size:fontSize})}); alert('弹窗样式已更新'); }
async function setPopupMaxHeight() { let height = parseInt(document.getElementById('popupMaxHeight').value); await fetch('/set_popup_max_height', {method:'POST', headers:{'Content-Type':'application/json'}, body:JSON.stringify({height:height})}); alert('弹窗最大高度已更新'); }
async function setPopupAutoScroll() { let enabled = document.getElementById('popupAutoScroll').checked; await fetch('/set_popup_auto_scroll', {method:'POST', headers:{'Content-Type':'application/json'}, body:JSON.stringify({enabled:enabled})}); alert('自动滚动设置已保存'); }
async function setPopupScrollSpeed() { let speed = parseInt(document.getElementById('popupVerticalScrollSpeed').value); await fetch('/set_popup_scroll_speed', {method:'POST', headers:{'Content-Type':'application/json'}, body:JSON.stringify({speed:speed})}); document.getElementById('popupVerticalScrollSpeedVal').innerText = speed + 'ms'; alert('竖向滚动速度已更新'); }
document.getElementById('popupVerticalScrollSpeed').oninput = function() { document.getElementById('popupVerticalScrollSpeedVal').innerText = this.value + 'ms'; }

function openChangePasswordModal() {
    document.getElementById('changePwdModal').style.display = 'flex';
    document.getElementById('oldPwd').value = '';
    document.getElementById('newPwd').value = '';
    document.getElementById('confirmPwd').value = '';
}
async function changePassword() {
    let type = document.getElementById('pwdType').value;
    let oldPwd = document.getElementById('oldPwd').value;
    let newPwd = document.getElementById('newPwd').value;
    let confirmPwd = document.getElementById('confirmPwd').value;
    if (newPwd !== confirmPwd) {
        alert('两次输入的新密码不一致');
        return;
    }
    if (newPwd.length < 4) {
        alert('新密码长度至少4位');
        return;
    }
    let url = type === 'web' ? '/change_web_password' : '/change_exit_password';
    let response = await fetch(url, {
        method: 'POST',
        headers: {'Content-Type': 'application/json'},
        body: JSON.stringify({old_password: oldPwd, new_password: newPwd})
    });
    let data = await response.json();
    if (data.error) {
        alert('修改失败: ' + data.error);
    } else {
        alert('修改成功！');
        closeModal('changePwdModal');
    }
}

// 名单导入
async function openRosterModal() { document.getElementById('rosterModal').style.display='flex'; document.getElementById('rosterFile').value=''; document.getElementById('rosterStatus').innerHTML=''; document.getElementById('searchName').value=''; allNames=[]; currentNames=[]; document.getElementById('nameListContainer').innerHTML=''; }
function closeRosterModal() { document.getElementById('rosterModal').style.display='none'; }
async function uploadRoster() { let fileInput=document.getElementById('rosterFile'); let file=fileInput.files[0]; if(!file){alert('请选择文件'); return;} let formData=new FormData(); formData.append('file',file); let statusDiv=document.getElementById('rosterStatus'); statusDiv.innerHTML='上传中...'; try{ let response=await fetch('/upload_roster',{method:'POST',body:formData}); let data=await response.json(); if(data.error) statusDiv.innerHTML='<span style="color:red;">错误: '+data.error+'</span>'; else{ allNames=data.names; currentNames=[...allNames]; statusDiv.innerHTML='<span style="color:green;">成功导入 '+allNames.length+' 个名字</span>'; renderNameList(currentNames); } }catch(e){ statusDiv.innerHTML='<span style="color:red;">请求失败</span>'; } }
function renderNameList(names){ let container=document.getElementById('nameListContainer'); if(!names.length){ container.innerHTML='<div style="padding:10px;text-align:center;">暂无数据</div>'; return; } let html=''; for(let name of names){ html+=`<div class="name-item" onclick="insertName('${name.replace(/'/g, "\\'")}')">${escapeHtml(name)}</div>`; } container.innerHTML=html; }
function escapeHtml(str){ return str.replace(/[&<>]/g, function(m){ if(m==='&') return '&amp;'; if(m==='<') return '&lt;'; if(m==='>') return '&gt;'; return m;}); }
function filterNames(){ let keyword=document.getElementById('searchName').value.trim().toLowerCase(); if(!keyword) currentNames=[...allNames]; else currentNames=allNames.filter(n=>n.toLowerCase().includes(keyword)); renderNameList(currentNames); }
function insertName(name){ let textarea=document.getElementById('text'); let start=textarea.selectionStart; let end=textarea.selectionEnd; let value=textarea.value; let newValue=value.substring(0,start)+name+value.substring(end); textarea.value=newValue; textarea.focus(); textarea.setSelectionRange(start+name.length,start+name.length); closeRosterModal(); }

refreshQueue(); setInterval(refreshQueue,2000);
</script>
</body>
</html>"""

@app.route('/')
def index_redirect():
    if session.get('logged_in'):
        return redirect(url_for('dashboard'))
    return redirect(url_for('login_page'))

@app.route('/dashboard')
@login_required
def dashboard():
    return render_template_string(HTML_TEMPLATE, voices=voice_options, device_code=DEVICE_CODE,
                                  bg_color=bg_color, sender_name=sender_name,
                                  display_mode=display_mode, speech_enabled=speech_enabled,
                                  banner_bg_color=banner_bg_color, banner_text_color=banner_text_color,
                                  banner_height=banner_height, banner_font_size=banner_font_size,
                                  banner_scroll_speed=banner_scroll_speed,
                                  classic_alpha=classic_alpha, banner_alpha=banner_alpha,
                                  classic_text_color=text_color,
                                  popup_bg_color=popup_bg_color, popup_text_color=popup_text_color,
                                  popup_font_size=popup_font_size, popup_max_height=popup_max_height,
                                  popup_auto_scroll=popup_auto_scroll, popup_vertical_scroll_speed=popup_vertical_scroll_speed)

# ---------- API 路由（需登录并记录用户操作）----------
@app.route('/add_task', methods=['POST'])
@api_login_required
def add_task():
    record_user_action()
    d = request.get_json()
    task = {"type": d.get('type', 'normal'), "text": d.get('text', ''),
            "voice": d.get('voice', 'zh-CN-XiaoxiaoNeural'), "volume": d.get('volume', 1000),
            "always_show": d.get('always_show', False), "show_duration": int(d.get('show_duration', 0))}
    if task['type'] == 'repeat':
        task['repeat_count'] = int(d.get('repeat_count', 1))
        task['interval'] = int(d.get('interval', 0))
    elif task['type'] == 'timed':
        ts = d.get('trigger_time')
        if ts:
            task['trigger_time'] = datetime.datetime.fromisoformat(ts)
    queue_mgr.add(task)
    return {'status': 'ok'}

@app.route('/queue_status')
@api_login_required
def queue_status():
    # 轮询不记录用户操作，避免频繁更新在线时间
    s = queue_mgr.get_status()
    tasks = []
    for t in s['tasks']:
        t2 = dict(t)
        if 'trigger_time' in t2 and isinstance(t2['trigger_time'], datetime.datetime):
            t2['trigger_time'] = t2['trigger_time'].isoformat()
        tasks.append(t2)
    return {"tasks": tasks, "mode": s["mode"], "interval": s["interval"], "queue_count": s["queue_count"]}

@app.route('/set_mode', methods=['POST'])
@api_login_required
def set_mode():
    record_user_action()
    d = request.get_json()
    queue_mgr.set_mode(d['mode'], int(d.get('interval', 2)))
    return {'status': 'ok'}

@app.route('/manual_next', methods=['POST'])
@api_login_required
def manual_next():
    record_user_action()
    queue_mgr.manual_next()
    return {'status': 'ok'}

@app.route('/pause_queue', methods=['POST'])
@api_login_required
def pause_queue():
    record_user_action()
    queue_mgr.pause()
    return {'status': 'ok'}

@app.route('/clear_queue', methods=['POST'])
@api_login_required
def clear_queue():
    record_user_action()
    queue_mgr.clear()
    return {'status': 'ok'}

@app.route('/remove_task', methods=['POST'])
@api_login_required
def remove_task():
    record_user_action()
    queue_mgr.remove(request.get_json()['index'])
    return {'status': 'ok'}

@app.route('/move_task', methods=['POST'])
@api_login_required
def move_task():
    record_user_action()
    d = request.get_json()
    queue_mgr.move(d['index'], d['direction'])
    return {'status': 'ok'}

@app.route('/stop', methods=['POST'])
@api_login_required
def stop_all():
    record_user_action()
    queue_mgr.clear()
    return {'status': 'ok'}

@app.route('/set_sender', methods=['POST'])
@api_login_required
def set_sender():
    record_user_action()
    global sender_name, config
    name = request.get_json().get('name', '班呼ClassCall')
    sender_name = name
    config['sender_name'] = name
    save_config(config)
    if classic_window:
        root.after(0, lambda: classic_window.update_sender(name))
    return {'status': 'ok'}

@app.route('/toggle_green', methods=['POST'])
@api_login_required
def toggle_green():
    record_user_action()
    if classic_window:
        root.after(0, classic_window.toggle_green)
    return {'status': 'ok'}

@app.route('/toggle_fullscreen', methods=['POST'])
@api_login_required
def toggle_fullscreen():
    record_user_action()
    if classic_window:
        root.after(0, classic_window.toggle_fullscreen)
    return {'status': 'ok'}

@app.route('/set_bg_color', methods=['POST'])
@api_login_required
def set_bg_color():
    record_user_action()
    global bg_color, config
    color = request.get_json().get('color', 'black')
    bg_color = color
    config['bg_color'] = color
    save_config(config)
    if classic_window:
        root.after(0, lambda: classic_window._apply_bg(color) if not classic_window.green_mode else None)
    return {'status': 'ok'}

@app.route('/set_font', methods=['POST'])
@api_login_required
def set_font():
    record_user_action()
    global text_color, config
    d = request.get_json()
    family = d.get('family')
    size = d.get('size')
    color = d.get('color')
    if color:
        text_color = color
        config['text_color'] = color
        save_config(config)
    if classic_window:
        root.after(0, lambda: classic_window.update_font(family, size, color))
    return {'status': 'ok'}

@app.route('/set_display_mode', methods=['POST'])
@api_login_required
def set_display_mode():
    record_user_action()
    global display_mode, config
    mode = request.get_json().get('mode')
    if mode in ['classic', 'banner', 'notification']:
        display_mode = mode
        config['display_mode'] = mode
        save_config(config)
        if mode == 'classic' and classic_window:
            classic_window.clear_text()
        elif mode != 'classic' and classic_window:
            classic_window.hide_window(animated=False)
    return {'status': 'ok'}

@app.route('/set_speech_enabled', methods=['POST'])
@api_login_required
def set_speech_enabled():
    record_user_action()
    global speech_enabled, config
    enabled = request.get_json().get('enabled', False)
    speech_enabled = enabled
    config['speech_enabled'] = enabled
    save_config(config)
    return {'status': 'ok'}

@app.route('/set_classic_alpha', methods=['POST'])
@api_login_required
def set_classic_alpha():
    record_user_action()
    global classic_alpha, config
    alpha = request.get_json().get('alpha', 1.0)
    classic_alpha = max(0.0, min(1.0, alpha))
    config['classic_alpha'] = classic_alpha
    save_config(config)
    if classic_window:
        classic_window.set_alpha(classic_alpha)
    return {'status': 'ok'}

@app.route('/set_banner_alpha', methods=['POST'])
@api_login_required
def set_banner_alpha():
    record_user_action()
    global banner_alpha, config, banner_window
    alpha = request.get_json().get('alpha', 1.0)
    banner_alpha = max(0.0, min(1.0, alpha))
    config['banner_alpha'] = banner_alpha
    save_config(config)
    if banner_window:
        banner_window.set_alpha(banner_alpha)
    return {'status': 'ok'}

@app.route('/set_banner_scroll_speed', methods=['POST'])
@api_login_required
def set_banner_scroll_speed():
    record_user_action()
    global banner_scroll_speed, config, banner_window
    speed = request.get_json().get('speed', 100)
    banner_scroll_speed = max(30, min(500, speed))
    config['banner_scroll_speed'] = banner_scroll_speed
    save_config(config)
    if banner_window:
        banner_window.set_scroll_speed(banner_scroll_speed)
    return {'status': 'ok'}

@app.route('/set_banner_colors', methods=['POST'])
@api_login_required
def set_banner_colors():
    record_user_action()
    global banner_bg_color, banner_text_color, config, banner_window
    data = request.get_json()
    banner_bg_color = data.get('bg', '#2c3e50')
    banner_text_color = data.get('fg', 'white')
    config['banner_bg_color'] = banner_bg_color
    config['banner_text_color'] = banner_text_color
    save_config(config)
    if banner_window:
        banner_window.set_bg_color(banner_bg_color)
        banner_window.set_fg_color(banner_text_color)
    return {'status': 'ok'}

@app.route('/set_banner_style', methods=['POST'])
@api_login_required
def set_banner_style():
    record_user_action()
    global banner_height, banner_font_size, config, banner_window
    data = request.get_json()
    if 'height' in data:
        banner_height = data['height']
        config['banner_height'] = banner_height
    if 'font_size' in data:
        banner_font_size = data['font_size']
        config['banner_font_size'] = banner_font_size
    save_config(config)
    if banner_window:
        banner_window.set_height(banner_height)
        banner_window.set_font_size(banner_font_size)
    return {'status': 'ok'}

@app.route('/set_popup_bg_color', methods=['POST'])
@api_login_required
def set_popup_bg_color():
    record_user_action()
    global popup_bg_color, config, popup_notification
    color = request.get_json().get('color', '#2c3e50')
    popup_bg_color = color
    config['popup_bg_color'] = color
    save_config(config)
    if popup_notification:
        popup_notification.set_bg_color(color)
    return {'status': 'ok'}

@app.route('/set_popup_text_color', methods=['POST'])
@api_login_required
def set_popup_text_color():
    record_user_action()
    global popup_text_color, config, popup_notification
    color = request.get_json().get('color', 'white')
    popup_text_color = color
    config['popup_text_color'] = color
    save_config(config)
    if popup_notification:
        popup_notification.set_text_color(color)
    return {'status': 'ok'}

@app.route('/set_popup_style', methods=['POST'])
@api_login_required
def set_popup_style():
    record_user_action()
    global popup_font_size, config, popup_notification
    data = request.get_json()
    if 'font_size' in data:
        popup_font_size = data['font_size']
        config['popup_font_size'] = popup_font_size
        save_config(config)
        if popup_notification:
            popup_notification.set_font_size(popup_font_size)
    return {'status': 'ok'}

@app.route('/set_popup_max_height', methods=['POST'])
@api_login_required
def set_popup_max_height():
    record_user_action()
    global popup_max_height, config, popup_notification
    height = request.get_json().get('height', 300)
    popup_max_height = max(100, min(600, height))
    config['popup_max_height'] = popup_max_height
    save_config(config)
    if popup_notification:
        popup_notification.set_max_height(popup_max_height)
    return {'status': 'ok'}

@app.route('/set_popup_auto_scroll', methods=['POST'])
@api_login_required
def set_popup_auto_scroll():
    record_user_action()
    global popup_auto_scroll, config, popup_notification
    enabled = request.get_json().get('enabled', False)
    popup_auto_scroll = enabled
    config['popup_auto_scroll'] = popup_auto_scroll
    save_config(config)
    if popup_notification:
        popup_notification.set_auto_scroll(popup_auto_scroll)
    return {'status': 'ok'}

@app.route('/set_popup_scroll_speed', methods=['POST'])
@api_login_required
def set_popup_scroll_speed():
    record_user_action()
    global popup_vertical_scroll_speed, config, popup_notification
    speed = request.get_json().get('speed', 50)
    popup_vertical_scroll_speed = max(20, min(500, speed))
    config['popup_vertical_scroll_speed'] = popup_vertical_scroll_speed
    save_config(config)
    if popup_notification:
        popup_notification.set_scroll_speed(popup_vertical_scroll_speed)
    return {'status': 'ok'}

@app.route('/api/online_status')
@api_login_required
def online_status():
    global last_user_action_time
    is_online = False
    if last_user_action_time:
        diff = (datetime.datetime.now() - last_user_action_time).total_seconds()
        if diff < 60:
            is_online = True
    return jsonify({"online": is_online, "last_seen": last_user_action_time.isoformat() if last_user_action_time else None})

# ================== 接收端主界面 (托盘打开，功能完善) ==================
class MainUI:
    def __init__(self, root, device_code):
        self.root = root
        self.device_code = device_code
        self.window = None
        self.online_label = None
        self.device_label = None
        self.update_job = None

    def show(self):
        if self.window and self.window.winfo_exists():
            # 如果窗口已被 withdraw，则重新显示
            if not self.window.winfo_viewable():
                self.window.deiconify()
            self.window.lift()
            self.window.focus_force()
            return
        self.window = tk.Toplevel(self.root)
        self.window.title("班呼ClassCall - 接收端控制台")
        self.window.geometry("500x480")
        self.window.minsize(400, 400)
        self.window.configure(bg="#f7f9fc")
        self.window.protocol("WM_DELETE_WINDOW", self.hide)
        
        main_frame = tk.Frame(self.window, bg="#f7f9fc")
        main_frame.pack(fill=tk.BOTH, expand=True, padx=20, pady=20)
        
        title = tk.Label(main_frame, text="📢 班呼ClassCall 接收端", font=("微软雅黑", 18, "bold"), bg="#f7f9fc", fg="#2d3748")
        title.pack(pady=(0, 20))
        
        device_frame = tk.Frame(main_frame, bg="white", relief=tk.RIDGE, bd=0, highlightthickness=0)
        device_frame.pack(fill=tk.X, pady=10)
        device_frame.configure(height=80)
        tk.Label(device_frame, text="🔑 设备码", font=("微软雅黑", 12), bg="white", fg="#718096").pack(anchor="w", padx=15, pady=(10,0))
        self.device_label = tk.Label(device_frame, text=self.device_code, font=("Consolas", 20, "bold"), bg="white", fg="#4a5568")
        self.device_label.pack(anchor="w", padx=15, pady=(0,10))
        
        status_frame = tk.Frame(main_frame, bg="white", relief=tk.RIDGE, bd=0)
        status_frame.pack(fill=tk.X, pady=10)
        status_frame.configure(height=80)
        tk.Label(status_frame, text="📡 发送端在线状态", font=("微软雅黑", 12), bg="white", fg="#718096").pack(anchor="w", padx=15, pady=(10,0))
        self.online_label = tk.Label(status_frame, text="检测中...", font=("微软雅黑", 14, "bold"), bg="white", fg="#e53e3e")
        self.online_label.pack(anchor="w", padx=15, pady=(0,10))
        
        # 功能按钮区域
        func_frame = tk.Frame(main_frame, bg="#f7f9fc")
        func_frame.pack(fill=tk.X, pady=10)
        tk.Label(func_frame, text="🔧 快速操作", font=("微软雅黑", 12, "bold"), bg="#f7f9fc", fg="#4a5568").pack(anchor="w", pady=(0,5))
        btn_frame1 = tk.Frame(func_frame, bg="#f7f9fc")
        btn_frame1.pack(fill=tk.X)
        ttk.Button(btn_frame1, text="📱 显示二维码", command=self.show_qr).pack(side=tk.LEFT, padx=5, expand=True, fill=tk.X)
        ttk.Button(btn_frame1, text="🌐 打开控制面板", command=self.open_web).pack(side=tk.LEFT, padx=5, expand=True, fill=tk.X)
        
        btn_frame2 = tk.Frame(func_frame, bg="#f7f9fc")
        btn_frame2.pack(fill=tk.X, pady=(5,0))
        ttk.Button(btn_frame2, text="🔒 锁定界面", command=self.lock_ui).pack(side=tk.LEFT, padx=5, expand=True, fill=tk.X)
        ttk.Button(btn_frame2, text="🚪 退出程序", command=self.exit_program).pack(side=tk.LEFT, padx=5, expand=True, fill=tk.X)
        
        info_text = "💡 提示：\n• 锁定界面仅隐藏本窗口，程序仍在后台运行\n• 退出程序需输入密码（默认123456）\n• 控制面板地址将自动打开浏览器"
        info_label = tk.Label(main_frame, text=info_text, font=("微软雅黑", 9), bg="#f7f9fc", fg="#a0aec0", justify=tk.LEFT)
        info_label.pack(fill=tk.X, pady=(15,0))
        
        self.start_updating()
        self.window.update_idletasks()
        x = (self.window.winfo_screenwidth() // 2) - (self.window.winfo_width() // 2)
        y = (self.window.winfo_screenheight() // 2) - (self.window.winfo_height() // 2)
        self.window.geometry(f"+{x}+{y}")
        self.window.deiconify()
    
    def hide(self):
        if self.window:
            self.window.withdraw()
    
    def show_qr(self):
        ip = get_local_ip()
        port = 5000
        url = f"http://{ip}:{port}"
        qr_path = save_qr_image(url)
        try:
            webbrowser.open(qr_path)
        except:
            messagebox.showinfo("二维码", f"二维码已保存至: {qr_path}\n控制台地址: {url}")
    
    def open_web(self):
        ip = get_local_ip()
        port = 5000
        url = f"http://{ip}:{port}"
        webbrowser.open(url)
    
    def lock_ui(self):
        # 仅隐藏窗口，不退出程序
        self.hide()
    
    def exit_program(self):
        pwd = simpledialog.askstring("密码验证", "请输入退出密码:", show='*')
        if pwd == EXIT_PASSWORD:
            self.quit_app()
        else:
            messagebox.showerror("错误", "密码错误，无法退出")
    
    def quit_app(self):
        debug("用户通过托盘退出程序")
        if queue_mgr:
            queue_mgr.running = False
        if tray_icon:
            tray_icon.stop()
        root.after(100, root.destroy)
        os._exit(0)
    
    def start_updating(self):
        self.update_online_status()
        self.update_job = self.window.after(3000, self.start_updating)
    
    def update_online_status(self):
        global last_user_action_time
        try:
            if last_user_action_time:
                diff = (datetime.datetime.now() - last_user_action_time).total_seconds()
                if diff < 60:
                    self.online_label.config(text="● 在线 (最近有操作)", fg="#38a169")
                else:
                    self.online_label.config(text="○ 离线 (超过1分钟无操作)", fg="#e53e3e")
            else:
                self.online_label.config(text="○ 未检测到连接", fg="#e53e3e")
        except:
            pass

# ================== 托盘图标 ==================
def setup_tray():
    global tray_icon, main_ui_window
    try:
        image = Image.new('RGB', (64, 64), color='#667eea')
        draw = ImageDraw.Draw(image)
        draw.rectangle([16,16,48,48], fill='white')
        draw.text((24,24), "📢", fill='#667eea')
    except:
        image = Image.open(os.path.join(BASE_DIR, "icon.png")) if os.path.exists(os.path.join(BASE_DIR, "icon.png")) else None
    if image is None:
        image = Image.new('RGB', (64, 64), color='#667eea')
    menu = pystray.Menu(
        pystray.MenuItem("打开主界面", lambda: root.after(0, lambda: main_ui_window.show())),
        pystray.MenuItem("退出程序", lambda: root.after(0, lambda: main_ui_window.exit_program()))
    )
    tray_icon = pystray.Icon("BanhuClassCall", image, "班呼ClassCall", menu)
    tray_icon.run_detached()

# ================== 网络工具 ==================
def get_local_ip():
    try:
        s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        s.connect(("8.8.8.8", 80))
        ip = s.getsockname()[0]
        s.close()
        return ip
    except:
        return "127.0.0.1"

def save_qr_image(url):
    target_dir = os.path.join(BASE_DIR, "qr_codes")
    try:
        os.makedirs(target_dir, exist_ok=True)
        path = os.path.join(target_dir, "qr_code.png")
        qrcode.make(url).save(path)
        return path
    except:
        fallback_dir = os.path.join(tempfile.gettempdir(), "call_device_qr")
        os.makedirs(fallback_dir, exist_ok=True)
        path = os.path.join(fallback_dir, "qr_code.png")
        qrcode.make(url).save(path)
        return path

# ================== 启动 ==================
def main():
    global classic_window, banner_window, root, popup_notification, queue_mgr, main_ui_window, last_user_action_time
    port = 5000
    ip = get_local_ip()
    url = f"http://{ip}:{port}"
    qr_path = save_qr_image(url)
    debug(f"二维码已保存至: {qr_path}")
    debug(f"控制面板地址: {url} (密码: {WEB_PASSWORD})")
    
    root = tk.Tk()
    root.withdraw()
    
    # 接收端主界面
    main_ui_window = MainUI(root, DEVICE_CODE)
    # 启动后自动显示配对二维码
    root.after(500, lambda: main_ui_window.show_qr())
    # 自动打开主界面
    root.after(1000, lambda: main_ui_window.show())
    
    # 初始化显示组件
    classic_window = ClassicWindow(root, DEVICE_CODE, bg_color, sender_name)
    if display_mode != 'classic':
        classic_window.hide_window(animated=False)
    
    banner_window = BannerWindow()
    banner_window.set_bg_color(banner_bg_color)
    banner_window.set_fg_color(banner_text_color)
    banner_window.set_height(banner_height)
    banner_window.set_font_size(banner_font_size)
    banner_window.set_alpha(banner_alpha)
    banner_window.set_scroll_speed(banner_scroll_speed)
    
    popup_notification = PopupNotification(root)
    popup_notification.set_bg_color(popup_bg_color)
    popup_notification.set_text_color(popup_text_color)
    popup_notification.set_font_size(popup_font_size)
    popup_notification.set_max_height(popup_max_height)
    popup_notification.set_auto_scroll(popup_auto_scroll)
    popup_notification.set_scroll_speed(popup_vertical_scroll_speed)
    
    queue_mgr = TaskQueue()
    last_user_action_time = None
    
    # 启动Flask线程
    threading.Thread(target=app.run, kwargs={'host': '0.0.0.0', 'port': port, 'debug': False, 'use_reloader': False}, daemon=True).start()
    
    # 启动托盘
    setup_tray()
    
    root.mainloop()

if __name__ == '__main__':
    main()